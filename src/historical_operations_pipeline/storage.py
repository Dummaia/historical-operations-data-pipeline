"""Persistence adapters for files and PostgreSQL."""

from __future__ import annotations

import json
import re
import tempfile
from pathlib import Path
from typing import TYPE_CHECKING

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from historical_operations_pipeline.config import PipelineConfig
from historical_operations_pipeline.exceptions import PipelineError

if TYPE_CHECKING:
    from historical_operations_pipeline.models import PipelineResult

HEADER_FILL = PatternFill(fill_type="solid", fgColor="172554")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SEVERITY_FILLS = {
    "error": PatternFill(fill_type="solid", fgColor="FEE2E2"),
    "warning": PatternFill(fill_type="solid", fgColor="FEF3C7"),
    "info": PatternFill(fill_type="solid", fgColor="DBEAFE"),
}


def _temporary_path(directory: Path, suffix: str) -> Path:
    with tempfile.NamedTemporaryFile(
        dir=directory,
        prefix=".pipeline-",
        suffix=suffix,
        delete=False,
    ) as handle:
        return Path(handle.name)


def _format_excel(path: Path) -> None:
    workbook = load_workbook(path)
    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"
        worksheet.auto_filter.ref = worksheet.dimensions
        worksheet.sheet_view.zoomScale = 85
        for cell in worksheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )
        worksheet.row_dimensions[1].height = 32

        for column_cells in worksheet.iter_cols():
            header = str(column_cells[0].value or "")
            values = [str(cell.value) for cell in column_cells[:250] if cell.value is not None]
            width = min(max((len(value) for value in values), default=10) + 2, 42)
            width = min(max(width, len(header) + 4), 42)
            if header.endswith("_DATE"):
                width = 18
                for cell in column_cells[1:]:
                    cell.number_format = "yyyy-mm-dd"
            elif header in {
                "SNAPSHOT_ORDER",
                "SHEET_ORDER",
                "SOURCE_ROW_NUMBER",
                "SLA_DAYS_REMAINING",
                "FOLLOW_UP_AGE_DAYS",
                "RECORD_COUNT",
                "UNIQUE_PROCESSES",
                "count",
            }:
                for cell in column_cells[1:]:
                    cell.number_format = "#,##0"
            elif header == "severity":
                for cell in column_cells[1:]:
                    fill = SEVERITY_FILLS.get(str(cell.value).lower())
                    if fill is not None:
                        cell.fill = fill
                        cell.font = Font(bold=True)
            worksheet.column_dimensions[column_cells[0].column_letter].width = width

    workbook.save(path)


def write_outputs(
    result: PipelineResult,
    *,
    output_dir: str | Path,
    config: PipelineConfig,
) -> dict[str, Path]:
    """Atomically write Parquet, formatted Excel and JSON quality outputs."""

    directory = Path(output_dir)
    directory.mkdir(parents=True, exist_ok=True)

    parquet_path = directory / config.parquet_filename
    parquet_temp = _temporary_path(directory, ".parquet")
    try:
        result.history.to_parquet(parquet_temp, index=False)
        parquet_temp.replace(parquet_path)
    finally:
        parquet_temp.unlink(missing_ok=True)

    excel_path = directory / config.excel_filename
    excel_temp = _temporary_path(directory, ".xlsx")
    try:
        with pd.ExcelWriter(excel_temp, engine="openpyxl") as writer:
            result.history.to_excel(writer, sheet_name="history", index=False)
            result.latest_position.to_excel(
                writer,
                sheet_name="latest_position",
                index=False,
            )
            result.summary.to_excel(writer, sheet_name="summary", index=False)
            result.quality_report.issues_frame().to_excel(
                writer,
                sheet_name="quality_issues",
                index=False,
            )
        _format_excel(excel_temp)
        excel_temp.replace(excel_path)
    finally:
        excel_temp.unlink(missing_ok=True)

    quality_path = directory / config.quality_filename
    quality_temp = _temporary_path(directory, ".json")
    try:
        quality_temp.write_text(
            json.dumps(result.quality_report.to_dict(), indent=2, ensure_ascii=False),
            encoding="utf-8",
        )
        quality_temp.replace(quality_path)
    finally:
        quality_temp.unlink(missing_ok=True)

    paths = {
        "parquet": parquet_path,
        "excel": excel_path,
        "quality": quality_path,
    }
    result.output_paths.update(paths)
    return paths


def _safe_identifier(value: str) -> str:
    if not re.fullmatch(r"[A-Za-z_][A-Za-z0-9_]*", value):
        raise PipelineError(f"Unsafe PostgreSQL identifier: {value!r}")
    return value


def write_postgres(
    history: pd.DataFrame,
    *,
    database_url: str,
    config: PipelineConfig,
) -> None:
    """Persist the historical dataset using SQLAlchemy and chunked inserts."""

    try:
        from sqlalchemy import create_engine, text
    except ImportError as exc:  # pragma: no cover - optional dependency
        raise PipelineError(
            "PostgreSQL support is not installed. Run pip install -e '.[postgres]'."
        ) from exc

    schema = _safe_identifier(config.postgres_schema)
    table = _safe_identifier(config.postgres_table)
    database_frame = history.rename(columns=str.lower)
    engine = create_engine(database_url, pool_pre_ping=True)

    try:
        with engine.begin() as connection:
            connection.execute(text(f'CREATE SCHEMA IF NOT EXISTS "{schema}"'))
            database_frame.to_sql(
                table,
                connection,
                schema=schema,
                if_exists=config.postgres_if_exists,
                index=False,
                chunksize=config.postgres_chunksize,
                method="multi",
            )
    finally:
        engine.dispose()
